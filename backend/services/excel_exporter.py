import os
import re
from PIL import Image
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DISEASE_FULL_NAMES = {
    "ALL": "Acute Lymphoblastic Leukemia",
    "AML": "Acute Myeloid Leukemia",
    "CLL": "Chronic Lymphocytic Leukemia",
    "CML": "Chronic Myeloid Leukemia",
    "Normal": "Normal Non-Leukemic Smear",
    "NORMAL": "Normal Non-Leukemic Smear"
}


def _parse_dimensions(dim_str):
    """Parse width and height integers from dimension string like '4160 x 3120' or '4160 × 3120 pixels'."""
    if not dim_str or not isinstance(dim_str, str):
        return None, None
    matches = re.findall(r'\d+', dim_str)
    if len(matches) >= 2:
        try:
            return int(matches[0]), int(matches[1])
        except ValueError:
            pass
    return None, None


def _to_float(v, default=None):
    """Safely convert any numeric or string representation to a float."""
    if v is None:
        return default
    try:
        return float(v)
    except (ValueError, TypeError):
        return default


def _clean_conf(c_val):
    """Return numeric confidence or probability as-is without multiplying by 100."""
    v = _to_float(c_val)
    if v is None:
        return None
    return round(float(v), 1)


def generate_excel_report(analysis_data, output_path):
    """
    Generate a research evaluation Excel (.xlsx) report matching the exact visual specification:
    - Title banner at Row 1
    - Table header at Row 4 with Navy fill (#1E3A8A) and white bold text
    - 17 Columns including Final Fused Class Probabilities (ALL, AML, CLL, CML, Normal)
    - Dynamic data rows with proper numerical formatting (values added as-is without multiplying by 100)
    - AVERAGE / OVERALL SUMMARY row with soft ice-blue fill (#EEF5FE) and average metrics
    """
    # -------------------------------------------------------------------------
    # 1. Normalize Input Data
    # -------------------------------------------------------------------------
    if isinstance(analysis_data, list):
        items = analysis_data
    elif isinstance(analysis_data, dict):
        if "items" in analysis_data and isinstance(analysis_data["items"], list):
            items = analysis_data["items"]
        elif "results" in analysis_data and isinstance(analysis_data["results"], list):
            items = analysis_data["results"]
        else:
            items = [analysis_data]
    else:
        items = []

    # Filter out empty or non-dict items
    items = [it for it in items if isinstance(it, dict)]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Research Evaluation Report"

    # Explicitly show sheet grid lines
    ws.views.sheetView[0].showGridLines = True

    # -------------------------------------------------------------------------
    # 2. Typography & Color Palette
    # -------------------------------------------------------------------------
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=11, bold=True, color="000000")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=9.5, color="000000")
    summary_font = Font(name=font_family, size=9.5, bold=True, color="000000")

    # Header Fill: Exact Navy Blue (#1E3A8A)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    # Summary Fill: Soft Ice Blue (#EEF5FE)
    summary_fill = PatternFill(start_color="EEF5FE", end_color="EEF5FE", fill_type="solid")

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Borders
    border_cell = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    border_summary = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='B0C4DE'),
        bottom=Side(style='thin', color='B0C4DE')
    )

    # -------------------------------------------------------------------------
    # 3. Title Banner (Row 1) & Blank Spacers (Rows 2, 3)
    # -------------------------------------------------------------------------
    ws["A1"] = "LEUKEMIA DETECTION & RECONSTRUCTION - RESEARCH EVALUATION REPORT"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 8
    ws.row_dimensions[3].height = 8

    # -------------------------------------------------------------------------
    # 4. Table Headers (Row 4)
    # -------------------------------------------------------------------------
    columns = [
        {"header": "Slide #", "align": align_center, "width": 10},
        {"header": "Filename", "align": align_left, "width": 55},
        {"header": "Validation", "align": align_center, "width": 13},
        {"header": "Dimensions", "align": align_center, "width": 15},
        {"header": "Orig Size (KB)", "align": align_right, "width": 15, "fmt": "#,##0.0"},
        {"header": "Recon Size (KB)", "align": align_right, "width": 15, "fmt": "#,##0.0"},
        {"header": "PSNR (dB)", "align": align_right, "width": 12, "fmt": "0.00"},
        {"header": "SSIM", "align": align_right, "width": 12, "fmt": "0.0000"},
        {"header": "MAE", "align": align_right, "width": 12, "fmt": "0.0000"},
        {"header": "Hybrid Pred", "align": align_center, "width": 14},
        {"header": "Hybrid Conf (%)", "align": align_right, "width": 16, "fmt": '0.0"%"'},
        {"header": "ALL (%)", "align": align_right, "width": 11, "fmt": '0.0"%"'},
        {"header": "AML (%)", "align": align_right, "width": 11, "fmt": '0.0"%"'},
        {"header": "CLL (%)", "align": align_right, "width": 11, "fmt": '0.0"%"'},
        {"header": "CML (%)", "align": align_right, "width": 11, "fmt": '0.0"%"'},
        {"header": "Normal (%)", "align": align_right, "width": 13, "fmt": '0.0"%"'},
        {"header": "Full Disease Subtype", "align": align_left, "width": 32}
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 26
    for col_idx, col_info in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_info["header"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_cell

    # -------------------------------------------------------------------------
    # 5. Data Rows (Starting Row 5)
    # -------------------------------------------------------------------------
    valid_psnrs = []
    valid_ssims = []
    valid_maes = []
    valid_confs = []
    valid_alls = []
    valid_amls = []
    valid_clls = []
    valid_cmls = []
    valid_normals = []

    current_row = 5

    for idx, item in enumerate(items, start=1):
        # 1. Filename
        fname = (
            item.get("image_information", {}).get("filename") or
            item.get("selectedFileName") or
            item.get("filename") or
            f"Slide_{idx}.jpg"
        )

        # 2. Validation Status
        is_val = item.get("is_valid")
        if is_val is False:
            val_text = "Invalid"
        else:
            val_text = "Valid"

        # 3. Dynamic Dimensions
        dim_str = (
            item.get("image_information", {}).get("original_dimensions") or
            item.get("image_information", {}).get("dimensions") or
            item.get("original_dimensions") or
            item.get("dimensions") or ""
        )
        w, h = _parse_dimensions(dim_str)
        if not w and item.get("image_information", {}).get("width"):
            w = _to_float(item.get("image_information", {}).get("width"))
        if not h and item.get("image_information", {}).get("height"):
            h = _to_float(item.get("image_information", {}).get("height"))

        # If dimensions not pre-saved, inspect actual image file on disk if available
        if (not w or not h):
            orig_path = (
                item.get("image_paths", {}).get("original") or
                item.get("orig_path") or
                item.get("path")
            )
            if orig_path and os.path.exists(orig_path):
                try:
                    with Image.open(orig_path) as img:
                        w, h = img.size
                except Exception:
                    pass

        dim_display = f"{int(w)} x {int(h)}" if (w and h) else (dim_str if dim_str else "N/A")

        # 4. Dynamic File Sizes (KB)
        orig_bytes = _to_float(
            item.get("image_information", {}).get("original_file_size_bytes") or
            item.get("original_file_size_bytes") or
            item.get("original_size_bytes") or
            item.get("file_size_bytes")
        )
        if (orig_bytes is None or orig_bytes <= 0):
            orig_path = (
                item.get("image_paths", {}).get("original") or
                item.get("orig_path") or
                item.get("path")
            )
            if orig_path and os.path.exists(orig_path):
                try:
                    orig_bytes = float(os.path.getsize(orig_path))
                except Exception:
                    pass

        orig_kb = round(orig_bytes / 1024.0, 1) if (orig_bytes and orig_bytes > 0) else None

        recon_bytes = _to_float(
            item.get("image_information", {}).get("reconstructed_file_size_bytes") or
            item.get("reconstructed_file_size_bytes") or
            item.get("reconstructed_size_bytes") or
            item.get("recon_size_bytes") or
            item.get("recon_file_size_bytes")
        )
        if (recon_bytes is None or recon_bytes <= 0):
            recon_path = (
                item.get("image_paths", {}).get("reconstructed") or
                item.get("recon_path")
            )
            if recon_path and os.path.exists(recon_path):
                try:
                    recon_bytes = float(os.path.getsize(recon_path))
                except Exception:
                    pass

        recon_kb = round(recon_bytes / 1024.0, 1) if (recon_bytes and recon_bytes > 0) else None

        # 5. Dynamic Calculated Metrics (PSNR, SSIM, MAE)
        rm = item.get("reconstruction_metrics", {})
        psnr_val = _to_float(rm.get("psnr_db") if rm.get("psnr_db") is not None else rm.get("psnr"))
        if psnr_val is None:
            psnr_val = _to_float(item.get("psnr_db") if item.get("psnr_db") is not None else item.get("psnr"))

        ssim_val = _to_float(rm.get("ssim"))
        if ssim_val is None:
            ssim_val = _to_float(item.get("ssim"))

        mae_val = _to_float(rm.get("mae"))
        if mae_val is None:
            mae_val = _to_float(item.get("mae"))

        psnr_cell = round(psnr_val, 2) if psnr_val is not None else "N/A"
        ssim_cell = round(ssim_val, 4) if ssim_val is not None else "N/A"
        mae_cell = round(mae_val, 4) if mae_val is not None else "N/A"

        if psnr_val is not None and val_text == "Valid":
            valid_psnrs.append(psnr_val)
        if ssim_val is not None and val_text == "Valid":
            valid_ssims.append(ssim_val)
        if mae_val is not None and val_text == "Valid":
            valid_maes.append(mae_val)

        # 6. Hybrid Prediction & Confidence (as-is, do not multiply by 100)
        pred_cls = (
            item.get("hybrid_pred") or
            item.get("hybrid", {}).get("prediction") or
            item.get("classification", {}).get("prediction") or
            item.get("prediction") or
            item.get("resnet50_pred") or
            "N/A"
        )
        raw_conf = (
            item.get("hybrid_conf") or
            item.get("hybrid", {}).get("confidence") or
            item.get("classification", {}).get("confidence") or
            item.get("confidence")
        )
        conf_pct = _clean_conf(raw_conf)
        conf_cell = conf_pct if conf_pct is not None else "N/A"

        # 7. Final Fused Probabilities Breakdown (as-is, do not multiply by 100)
        probs = (
            item.get("hybrid", {}).get("probabilities") or
            item.get("classification", {}).get("probabilities") or
            item.get("probabilities") or {}
        )
        p_all = _clean_conf(probs.get("ALL"))
        p_aml = _clean_conf(probs.get("AML"))
        p_cll = _clean_conf(probs.get("CLL"))
        p_cml = _clean_conf(probs.get("CML"))
        p_normal = _clean_conf(probs.get("Normal"))

        if val_text == "Valid":
            if conf_pct is not None:
                valid_confs.append(conf_pct)
            if p_all is not None:
                valid_alls.append(p_all)
            if p_aml is not None:
                valid_amls.append(p_aml)
            if p_cll is not None:
                valid_clls.append(p_cll)
            if p_cml is not None:
                valid_cmls.append(p_cml)
            if p_normal is not None:
                valid_normals.append(p_normal)

        # 8. Full Disease Subtype
        full_subtype = DISEASE_FULL_NAMES.get(pred_cls, pred_cls)

        row_values = [
            idx,
            fname,
            val_text,
            dim_display,
            orig_kb if orig_kb is not None else "N/A",
            recon_kb if recon_kb is not None else "N/A",
            psnr_cell,
            ssim_cell,
            mae_cell,
            pred_cls,
            conf_cell,
            p_all if p_all is not None else "N/A",
            p_aml if p_aml is not None else "N/A",
            p_cll if p_cll is not None else "N/A",
            p_cml if p_cml is not None else "N/A",
            p_normal if p_normal is not None else "N/A",
            full_subtype
        ]

        ws.row_dimensions[current_row].height = 20
        for col_idx, (col_info, val) in enumerate(zip(columns, row_values), start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = col_info["align"]
            cell.border = border_cell

            if isinstance(val, (int, float)) and "fmt" in col_info:
                cell.number_format = col_info["fmt"]

        current_row += 1

    # -------------------------------------------------------------------------
    # 6. AVERAGE / OVERALL SUMMARY Row
    # -------------------------------------------------------------------------
    summary_row = current_row
    ws.row_dimensions[summary_row].height = 22

    # Pre-fill row cells with summary styling
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=summary_row, column=col_idx)
        cell.font = summary_font
        cell.fill = summary_fill
        cell.border = border_summary

    # Column B: AVERAGE / OVERALL SUMMARY label
    label_cell = ws.cell(row=summary_row, column=2, value="AVERAGE / OVERALL SUMMARY")
    label_cell.alignment = align_left

    # Dynamic average calculations
    if valid_psnrs:
        avg_psnr = round(sum(valid_psnrs) / len(valid_psnrs), 2)
        psnr_avg_cell = ws.cell(row=summary_row, column=7, value=avg_psnr)
        psnr_avg_cell.alignment = align_right
        psnr_avg_cell.number_format = "0.00"

    if valid_ssims:
        avg_ssim = round(sum(valid_ssims) / len(valid_ssims), 4)
        ssim_avg_cell = ws.cell(row=summary_row, column=8, value=avg_ssim)
        ssim_avg_cell.alignment = align_right
        ssim_avg_cell.number_format = "0.0000"

    if valid_maes:
        avg_mae = round(sum(valid_maes) / len(valid_maes), 4)
        mae_avg_cell = ws.cell(row=summary_row, column=9, value=avg_mae)
        mae_avg_cell.alignment = align_right
        mae_avg_cell.number_format = "0.0000"

    if valid_confs:
        avg_conf = round(sum(valid_confs) / len(valid_confs), 1)
        conf_avg_cell = ws.cell(row=summary_row, column=11, value=avg_conf)
        conf_avg_cell.alignment = align_right
        conf_avg_cell.number_format = '0.0"%"'

    if valid_alls:
        cell = ws.cell(row=summary_row, column=12, value=round(sum(valid_alls) / len(valid_alls), 1))
        cell.alignment = align_right
        cell.number_format = '0.0"%"'

    if valid_amls:
        cell = ws.cell(row=summary_row, column=13, value=round(sum(valid_amls) / len(valid_amls), 1))
        cell.alignment = align_right
        cell.number_format = '0.0"%"'

    if valid_clls:
        cell = ws.cell(row=summary_row, column=14, value=round(sum(valid_clls) / len(valid_clls), 1))
        cell.alignment = align_right
        cell.number_format = '0.0"%"'

    if valid_cmls:
        cell = ws.cell(row=summary_row, column=15, value=round(sum(valid_cmls) / len(valid_cmls), 1))
        cell.alignment = align_right
        cell.number_format = '0.0"%"'

    if valid_normals:
        cell = ws.cell(row=summary_row, column=16, value=round(sum(valid_normals) / len(valid_normals), 1))
        cell.alignment = align_right
        cell.number_format = '0.0"%"'

    # -------------------------------------------------------------------------
    # 7. Column Widths
    # -------------------------------------------------------------------------
    for col_idx, col_info in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        # Check max length of filenames or content to ensure no truncation
        calc_width = col_info["width"]
        if col_idx == 2 and items:  # Filename column
            max_len = max(len(str(ws.cell(row=r, column=2).value or '')) for r in range(5, current_row))
            calc_width = max(col_info["width"], min(max_len + 3, 110))
        ws.column_dimensions[col_letter].width = calc_width

    # Ensure parent directory exists and save
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    wb.save(output_path)
    return output_path
